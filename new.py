import os
import re
import openpyxl
import docx
from PyPDF2 import PdfReader

def extract_text_from_pdf(pdf_file):
    text = ""
    with open(pdf_file, 'rb') as file:
        reader = PdfReader(file)
        for page_num in range(len(reader.pages)):
            text += reader.pages[page_num].extract_text()
    return text

def extract_text_from_docx(docx_file):
    text = ""
    doc = docx.Document(docx_file)
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

def extract_text_from_doc(doc_file):
    text = ""
    with open(doc_file, 'rb') as file:
        # Read DOC file as binary
        doc_content = file.read()
        # Decode bytes to string using utf-8 encoding
        text = doc_content.decode("utf-8", errors="ignore")
    return text

def extract_contact_info_from_resume(file):
    contact_info = {}
    overall_text = ""
    if file.endswith('.pdf'):
        overall_text = extract_text_from_pdf(file)
    elif file.endswith('.docx'):
        overall_text = extract_text_from_docx(file)
    elif file.endswith('.doc'):
        overall_text = extract_text_from_doc(file)
    else:
        print(f"Unsupported file format: {file}")
        return contact_info, overall_text

    # Regular expression to extract email addresses
    email = re.findall(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', overall_text)
    # Regular expression to extract phone numbers
    phone = re.findall(r'\b(?:\+\d{1,2}\s)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}\b', overall_text)
    if email:
        contact_info['Email'] = email[0]
    if phone:
        contact_info['Phone'] = phone[0]
    return contact_info, overall_text

def main():
    resumes_directory = 'Sample2'
    output_excel_file = 'contact_info.xlsx'

    # Create a new Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    # Set column names
    ws.append(['Email', 'Phone', 'Overall Text'])

    # Loop through each file in the directory
    for filename in os.listdir(resumes_directory):
        file_path = os.path.join(resumes_directory, filename)
        if os.path.isfile(file_path):
            contact_info, overall_text = extract_contact_info_from_resume(file_path)
            if contact_info:
                # Write contact info and overall text to Excel file
                ws.append([contact_info.get('Email', ''), contact_info.get('Phone', ''), overall_text])

    # Save the Excel workbook
    wb.save(output_excel_file)

if __name__ == "__main__":
    main()
